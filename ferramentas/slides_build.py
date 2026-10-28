"""Gera a apresentação em HTML (padrão frontend-slides, palco fixo 1920×1080) e em PowerPoint editável.

Fonte única: conteudo/slides.json. Saídas: saida/apresentacao.html e saida/Masterclass-de-Hardware.pptx
"""
import html, io, json, os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, 'saida')
BASE_CSS = os.path.join(os.path.expanduser('~'), '.claude', 'skills', 'frontend-slides', 'viewport-base.css')
e = html.escape


# =================== HTML ===================

def chrome(s, i, total):
    k = s.get('k', '')
    return (f'<div class="sl-top"><span>{e(k)}</span><span class="sl-deck">Masterclass de Hardware</span></div>'
            f'<div class="sl-num">{i:02d}<i>/{total:02d}</i></div>')


def slide_html(s, i, total):
    t = s['tipo']
    if t == 'capa':
        return (f'<section class="slide escuro capa"><div class="grade"></div><div class="capa-c">'
                f'<p class="kick">{e(s["k"])}</p><h1 class="anim">{e(s["t"])}</h1><p class="sub anim d1">{e(s["s"])}</p></div>'
                f'<div class="capa-r anim d2"><span>Apostila · HTML interativo · Apresentação</span><span>{e(s["r"])}</span></div></section>')
    if t == 'secao':
        return (f'<section class="slide escuro secao"><div class="grade"></div><p class="sec-n anim">{e(s["n"])}</p>'
                f'<div class="sec-c"><h2 class="anim d1">{e(s["t"])}</h2><p class="anim d2">{e(s["s"])}</p></div>'
                f'<div class="sl-num claro">{i:02d}<i>/{total:02d}</i></div></section>')
    if t == 'fim':
        return (f'<section class="slide escuro capa fim"><div class="grade"></div><div class="capa-c">'
                f'<h1 class="anim">{e(s["t"])}</h1><p class="sub anim d1">{e(s["s"])}</p><p class="nota anim d2">{e(s["r"])}</p></div></section>')
    corpo = ''
    if t == 'cards':
        n = len(s['cards'])
        cols = 5 if n == 10 else 3
        itens = ''.join(
            f'<div class="card anim d{min(j, 5)}"><span class="c-k">{e(a)}</span><b>{e(b)}</b>{("<p>" + e(c) + "</p>") if c else ""}</div>'
            for j, (a, b, c) in enumerate(s['cards']))
        corpo = f'<div class="cards c{cols} n{n}">{itens}</div>'
    elif t == 'duas':
        def col(lista, cls):
            return (f'<div class="col {cls} anim"><h3>{e(lista[0])}</h3><p class="lead">{e(lista[1])}</p>'
                    f'<ul>{"".join("<li>" + e(x) + "</li>" for x in lista[2:])}</ul></div>')
        img = f'<figure class="duas-img anim d2"><img src="{s["img"]}" alt=""></figure>' if s.get('img') else ''
        corpo = f'<div class="duas{" com-img" if img else ""}">{col(s["a"], "ca")}{col(s["b"], "cb")}{img}</div>'
    elif t == 'tabela':
        cab = ''.join(f'<th>{e(c)}</th>' for c in s['cab'])
        linhas = ''.join('<tr>' + ''.join(f'<td>{e(c)}</td>' for c in l) + '</tr>' for l in s['linhas'])
        corpo = f'<table class="tab anim"><thead><tr>{cab}</tr></thead><tbody>{linhas}</tbody></table>'
    elif t == 'numeros':
        nums = ''.join(f'<div class="num anim d{min(j, 5)}"><b>{e(a)}</b><span>{e(b)}</span></div>' for j, (a, b) in enumerate(s['nums']))
        img = (f'<figure class="num-img anim d2"><img src="{s["img"]}" alt=""><figcaption>{e(s.get("legenda", ""))}</figcaption></figure>'
               if s.get('img') else '')
        corpo = f'<div class="nums n{len(s["nums"])}{" com-img" if img else ""}">{nums}</div>{img}'
    elif t == 'linha':
        itens = ''.join(f'<li class="anim d{min(j // 2, 5)}"><span class="l-a">{e(a)}</span><b>{e(b)}</b><p>{e(c)}</p></li>'
                        for j, (a, b, c) in enumerate(s['itens']))
        corpo = f'<ol class="linha">{itens}</ol>'
    elif t == 'fluxo':
        itens = ''.join(f'<div class="passo anim d{j}"><span class="p-n">{j + 1:02d}</span><b>{e(a)}</b><p>{e(b)}</p></div>'
                        for j, (a, b) in enumerate(s['passos']))
        corpo = f'<div class="fluxo">{itens}</div>'
    rodape = f'<p class="rodape anim d3">{e(s["rodape"])}</p>' if s.get('rodape') else ''
    return (f'<section class="slide claro-bg t-{t}">{chrome(s, i, total)}<div class="sl-c">'
            f'<h2 class="titulo anim">{e(s["t"])}</h2>{corpo}{rodape}</div></section>')


CSS = r'''
/* === TOKENS === */
:root{--stage-bg:#0B0C0E;--slide-bg:#F5F3EE;--tinta:#16181C;--tinta-2:#4A4F57;--cobre:#B4531F;--cobre-c:#E58A52;--pcb:#0F5A43;--linha:#D9D3C7;--papel-2:#ECE8DF}
/* === BASE === */
.slide{font-family:'IBM Plex Sans',sans-serif;color:var(--tinta)}
.escuro{background:#16181C;color:#F5F3EE}
.grade{position:absolute;inset:0;background:repeating-linear-gradient(0deg,transparent 0 120px,rgba(255,255,255,.045) 120px 121px),repeating-linear-gradient(90deg,transparent 0 120px,rgba(255,255,255,.045) 120px 121px)}
.grade::after{content:"";position:absolute;top:0;bottom:0;left:1640px;width:6px;background:var(--cobre)}
/* === CAPA / FIM === */
.capa-c{position:absolute;left:140px;top:190px;width:1440px}
.kick{font:600 26px 'IBM Plex Mono',monospace;letter-spacing:.08em;text-transform:uppercase;color:var(--cobre-c);margin:0 0 48px}
.capa h1{font:900 176px/.92 'Archivo',sans-serif;letter-spacing:-.04em;margin:0 0 48px}
.capa .sub{font-size:38px;line-height:1.35;color:#CFCBC2;margin:0;max-width:880px}
.capa-r{position:absolute;left:140px;right:140px;bottom:96px;border-top:3px solid var(--cobre-c);padding-top:28px;display:flex;justify-content:space-between;font:600 26px 'IBM Plex Mono',monospace;color:#CFCBC2}
.fim h1{font-size:200px}
.fim .nota{font-size:28px;color:#9A9EA6;margin-top:60px;max-width:900px;line-height:1.4}
/* === SEÇÃO === */
.secao .sec-n{position:absolute;left:140px;top:120px;font:900 360px/1 'Archivo',sans-serif;color:var(--cobre);margin:0;letter-spacing:-.04em}
.sec-c{position:absolute;left:140px;bottom:150px;width:1440px}
.secao h2{font:900 140px/.95 'Archivo',sans-serif;letter-spacing:-.035em;margin:0 0 28px}
.secao p{font-size:38px;color:#CFCBC2;margin:0}
/* === CHROME === */
.sl-top{position:absolute;left:120px;right:120px;top:64px;display:flex;justify-content:space-between;font:600 22px 'IBM Plex Mono',monospace;text-transform:uppercase;letter-spacing:.07em;color:var(--cobre)}
.sl-deck{color:#8A8F97}
.sl-num{position:absolute;right:120px;bottom:56px;font:600 22px 'IBM Plex Mono',monospace;color:var(--tinta)}
.sl-num i{font-style:normal;color:#8A8F97}
.sl-num.claro{color:#F5F3EE}
.sl-c{position:absolute;left:120px;right:120px;top:130px;bottom:120px;display:flex;flex-direction:column}
.titulo{font:900 84px/1 'Archivo',sans-serif;letter-spacing:-.03em;margin:0 0 56px;padding-bottom:28px;border-bottom:5px solid var(--tinta);max-width:1500px}
.rodape{margin:auto 0 0;padding-top:28px;font-size:31px;line-height:1.4;color:var(--tinta-2);max-width:1560px}
/* === CARDS === */
.cards{display:grid;gap:28px;flex:1;grid-auto-rows:1fr;align-content:stretch}
.cards.c3{grid-template-columns:repeat(3,1fr)}
.cards.c5{grid-template-columns:repeat(5,1fr)}
.card{background:#fff;border:2px solid var(--linha);border-radius:18px;padding:36px 38px;display:flex;flex-direction:column;gap:14px}
.c-k{font:700 26px 'IBM Plex Mono',monospace;color:var(--cobre)}
.card b{font:800 46px/1.1 'Archivo',sans-serif;letter-spacing:-.01em}
.card p{margin:0;font-size:31px;line-height:1.38;color:var(--tinta-2)}
.cards.n3 .card{padding:48px}.cards.n3 .card b{font-size:60px}.cards.n3 .card p{font-size:36px}
.cards.n10 .card b{font-size:38px}.cards.n10 .c-k{font-size:44px;font-family:'Archivo',sans-serif;font-weight:900}
.cards.n9{gap:20px}.cards.n9 .card b{font-size:36px}.cards.n9 .card{padding:18px 30px;gap:6px}.cards.n9 .card p{font-size:26px}.cards.n9 .c-k{font-size:22px}
/* === DUAS COLUNAS === */
.duas{display:grid;grid-template-columns:1fr 1fr;gap:40px;flex:1}
.duas.com-img{grid-template-columns:1fr 1fr 520px}
.col{border-radius:20px;padding:44px 48px;background:#fff;border:2px solid var(--linha)}
.col.cb{background:#16181C;color:#F5F3EE;border-color:#16181C}
.col h3{font:900 60px/1 'Archivo',sans-serif;margin:0 0 18px;letter-spacing:-.02em}
.col.ca h3{color:var(--pcb)}.col.cb h3{color:var(--cobre-c)}
.col .lead{font-size:36px;line-height:1.35;font-weight:600;margin:0 0 26px}
.col ul{list-style:none;margin:0;padding:0}
.col li{font-size:33px;line-height:1.38;padding:22px 0 22px 38px;border-top:2px solid var(--linha);position:relative}
.col.cb li{border-color:#33373D;color:#D7D3CB}
.col li::before{content:"";position:absolute;left:0;top:38px;width:14px;height:14px;border-radius:3px;background:var(--cobre)}
.duas-img{margin:0;border-radius:20px;overflow:hidden;border:2px solid var(--linha)}
.duas-img img{width:100%;height:100%;object-fit:cover;max-height:none}
/* === TABELA === */
.tab{width:100%;border-collapse:collapse;font-size:32px;line-height:1.3}
.tab th{background:#16181C;color:#F5F3EE;text-align:left;padding:24px 28px;font:700 23px 'IBM Plex Sans',sans-serif;text-transform:uppercase;letter-spacing:.05em}
.tab td{padding:24px 28px;border-bottom:2px solid var(--linha);background:#fff}
.tab tbody tr:nth-child(even) td{background:#FAF8F4}
.tab td:first-child{font-weight:700}
/* === NÚMEROS === */
.nums{display:grid;gap:24px;flex:1;align-content:center}
.nums.n3{grid-template-columns:repeat(3,1fr)}.nums.n4{grid-template-columns:repeat(4,1fr)}.nums.n5{grid-template-columns:repeat(5,1fr)}
.nums.com-img{grid-template-columns:1fr 1fr;width:880px;align-content:start}
.num{border-top:6px solid var(--cobre);padding-top:26px}
.num b{display:block;font:900 92px/1 'Archivo',sans-serif;letter-spacing:-.035em}
.nums.n5 .num b{font-size:76px}.nums.com-img .num b{font-size:70px}
.nums.n3 .num b{font-size:94px}
.num span{display:block;font-size:34px;color:var(--tinta-2);margin-top:14px}
.num-img{position:absolute;right:0;top:170px;width:700px;margin:0}
.num-img img{width:700px;height:430px;object-fit:cover;border-radius:18px;max-height:none;display:block}
.num-img figcaption{font-size:18px;color:#8A8F97;margin-top:10px}
.t-numeros .rodape{max-width:900px}
/* === LINHA DO TEMPO === */
.linha{list-style:none;margin:0;padding:0;display:grid;grid-template-columns:repeat(4,1fr);grid-template-rows:1fr 1fr;gap:56px 36px;position:relative;flex:1}
.linha li{border-left:5px solid var(--cobre);padding:0 0 0 24px}
.l-a{font:700 30px 'IBM Plex Mono',monospace;color:var(--cobre)}
.linha b{display:block;font:800 48px/1.1 'Archivo',sans-serif;margin:14px 0 10px}
.linha p{margin:0;font-size:31px;line-height:1.35;color:var(--tinta-2)}
/* === FLUXO === */
.fluxo{display:grid;grid-template-columns:repeat(3,1fr);gap:0;flex:1;align-content:center}
.passo{padding:52px 56px;border:2px solid var(--tinta);background:#fff;position:relative;min-height:380px}
.passo+.passo{border-left:0}
.passo:nth-child(2){background:#16181C;color:#F5F3EE}
.passo:nth-child(2) p{color:#CFCBC2}
.p-n{font:700 30px 'IBM Plex Mono',monospace;color:var(--cobre)}
.passo b{display:block;font:900 66px/1.05 'Archivo',sans-serif;margin:30px 0 18px;letter-spacing:-.02em}
.passo p{margin:0;font-size:36px;line-height:1.38;color:var(--tinta-2)}
/* === ANIMAÇÃO === */
.slide .anim{opacity:0;transform:translateY(24px);transition:opacity .55s cubic-bezier(.2,.7,.2,1),transform .55s cubic-bezier(.2,.7,.2,1)}
.slide.visible .anim{opacity:1;transform:none}
.slide.visible .d1{transition-delay:.08s}.slide.visible .d2{transition-delay:.16s}.slide.visible .d3{transition-delay:.24s}.slide.visible .d4{transition-delay:.32s}.slide.visible .d5{transition-delay:.4s}
@media print{.slide .anim{opacity:1!important;transform:none!important}}
/* === CONTROLES (fora do palco) === */
.deck-controls{display:flex;gap:10px;align-items:center;font:600 14px 'IBM Plex Mono',monospace;color:#CFCBC2;background:rgba(22,24,28,.82);padding:8px 10px;border-radius:12px;opacity:.0;transition:opacity .25s}
body:hover .deck-controls,.deck-controls:focus-within{opacity:1}
.deck-controls button{background:#2A2D33;color:#F5F3EE;border:0;border-radius:8px;width:38px;height:34px;font:inherit;font-size:18px;cursor:pointer}
.barra-prog{position:fixed;left:0;top:0;height:4px;background:#B4531F;z-index:1001;transition:width .3s}
[contenteditable="true"]{outline:2px dashed #E58A52;outline-offset:4px}
'''

JS = r'''
/* === CONTROLADOR DA APRESENTAÇÃO === */
class SlidePresentation {
  constructor() {
    this.slides = document.querySelectorAll('.slide');
    this.stage = document.getElementById('deckStage');
    this.atual = 0;
    const h = parseInt(location.hash.slice(1), 10);
    this.escala(); addEventListener('resize', () => this.escala());
    this.teclado(); this.toque(); this.roda(); this.edicao();
    document.getElementById('prev').onclick = () => this.ir(this.atual - 1);
    document.getElementById('next').onclick = () => this.ir(this.atual + 1);
    this.ir(isNaN(h) ? 0 : h - 1);
  }
  escala() {
    const f = Math.min(innerWidth / 1920, innerHeight / 1080);
    this.stage.style.transform = `translate(${(innerWidth - 1920 * f) / 2}px, ${(innerHeight - 1080 * f) / 2}px) scale(${f})`;
  }
  ir(i) {
    this.atual = Math.max(0, Math.min(i, this.slides.length - 1));
    this.slides.forEach((s, k) => { s.classList.toggle('active', k === this.atual); s.classList.toggle('visible', k === this.atual); });
    document.getElementById('cont').textContent = (this.atual + 1) + ' / ' + this.slides.length;
    document.getElementById('prog').style.width = (100 * (this.atual + 1) / this.slides.length) + '%';
    history.replaceState(null, '', '#' + (this.atual + 1));
  }
  teclado() {
    addEventListener('keydown', (e) => {
      if (document.activeElement && document.activeElement.isContentEditable) return;
      if (['ArrowRight', 'ArrowDown', 'PageDown', ' '].includes(e.key)) { e.preventDefault(); this.ir(this.atual + 1); }
      if (['ArrowLeft', 'ArrowUp', 'PageUp'].includes(e.key)) { e.preventDefault(); this.ir(this.atual - 1); }
      if (e.key === 'Home') this.ir(0);
      if (e.key === 'End') this.ir(this.slides.length - 1);
      if (e.key === 'f' || e.key === 'F') document.fullscreenElement ? document.exitFullscreen() : document.documentElement.requestFullscreen();
      if (e.key === 'e' || e.key === 'E') this.alternarEdicao();
    });
  }
  toque() {
    let x0 = null;
    addEventListener('touchstart', (e) => { x0 = e.touches[0].clientX; }, { passive: true });
    addEventListener('touchend', (e) => {
      if (x0 === null) return;
      const dx = e.changedTouches[0].clientX - x0;
      if (Math.abs(dx) > 40) this.ir(this.atual + (dx < 0 ? 1 : -1));
      x0 = null;
    });
  }
  roda() {
    let trava = false;
    addEventListener('wheel', (e) => {
      if (trava || Math.abs(e.deltaY) < 20) return;
      trava = true; this.ir(this.atual + (e.deltaY > 0 ? 1 : -1));
      setTimeout(() => { trava = false; }, 650);
    }, { passive: true });
  }
  /* === EDIÇÃO INLINE (tecla E) === */
  edicao() {
    this.editando = false;
    try {
      const salvo = JSON.parse(localStorage.getItem('mh-slides') || '{}');
      Object.keys(salvo).forEach((k) => { const el = document.querySelector('[data-ed="' + k + '"]'); if (el) el.innerHTML = salvo[k]; });
    } catch (e) {}
    document.getElementById('edit').onclick = () => this.alternarEdicao();
  }
  alternarEdicao() {
    this.editando = !this.editando;
    document.querySelectorAll('[data-ed]').forEach((el) => { el.contentEditable = this.editando ? 'true' : 'false'; });
    document.getElementById('edit').textContent = this.editando ? '✓' : '✎';
    if (!this.editando) {
      const dados = {};
      document.querySelectorAll('[data-ed]').forEach((el) => { dados[el.dataset.ed] = el.innerHTML; });
      try { localStorage.setItem('mh-slides', JSON.stringify(dados)); } catch (e) {}
    }
  }
}
new SlidePresentation();
'''


def gerar_html(slides):
    base = io.open(BASE_CSS, encoding='utf-8').read()
    total = len(slides)
    corpo = ''.join(slide_html(s, i + 1, total) for i, s in enumerate(slides))
    # marca os textos editáveis
    import re
    contador = [0]

    def marca(m):
        contador[0] += 1
        return f'<{m.group(1)} data-ed="t{contador[0]}"{m.group(2)}>'
    corpo = re.sub(r'<(h1|h2|h3|p|b|li|td|th)((?: class="[^"]*")?)>', marca, corpo)
    pagina = f'''<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Masterclass de Hardware — Apresentação</title>
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Archivo:wght@700;800;900&family=IBM+Plex+Sans:wght@400;600;700&family=IBM+Plex+Mono:wght@600;700&display=swap" rel="stylesheet">
<style>
{base}
@page{{size:1920px 1080px;margin:0}}
{CSS}
</style></head><body>
<div class="barra-prog" id="prog"></div>
<div class="deck-viewport"><main class="deck-stage" id="deckStage">
{corpo}
</main></div>
<div class="deck-controls"><button id="prev" aria-label="Slide anterior">‹</button><span id="cont"></span><button id="next" aria-label="Próximo slide">›</button><button id="edit" aria-label="Editar textos (tecla E)" title="Editar textos (E)">✎</button></div>
<script>{JS}</script>
</body></html>'''
    io.open(os.path.join(SAIDA, 'apresentacao.html'), 'w', encoding='utf-8').write(pagina)


# =================== PPTX ===================
W, H = Emu(12192000), Emu(6858000)  # 16:9
PX = W / 1920  # converte as medidas do palco HTML


def px(v):
    return Emu(int(v * PX))


TINTA = RGBColor(0x16, 0x18, 0x1C); PAPEL = RGBColor(0xF5, 0xF3, 0xEE); COBRE = RGBColor(0xB4, 0x53, 0x1F)
COBRE_C = RGBColor(0xE5, 0x8A, 0x52); CINZA = RGBColor(0x4A, 0x4F, 0x57); CINZA_C = RGBColor(0xCF, 0xCB, 0xC2)
LINHA = RGBColor(0xD9, 0xD3, 0xC7); BRANCO = RGBColor(0xFF, 0xFF, 0xFF); PCB = RGBColor(0x0F, 0x5A, 0x43)
F_TIT, F_TXT, F_MONO = 'Segoe UI Black', 'Segoe UI', 'Consolas'


def fundo(sl, cor):
    f = sl.background.fill; f.solid(); f.fore_color.rgb = cor


def caixa(sl, x, y, w, h, preench=None, borda=None, raio=False):
    forma = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if raio else MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    if raio:
        forma.adjustments[0] = 0.06
    if preench is None:
        forma.fill.background()
    else:
        forma.fill.solid(); forma.fill.fore_color.rgb = preench
    if borda is None:
        forma.line.fill.background()
    else:
        forma.line.color.rgb = borda; forma.line.width = Pt(1.25)
    forma.shadow.inherit = False
    return forma


def texto(sl, x, y, w, h, linhas, ancora=MSO_ANCHOR.TOP):
    """linhas: lista de (texto, tamanho_px, cor, fonte, negrito)."""
    tb = sl.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = ancora
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for k, (t, tam, cor, fonte, neg) in enumerate(linhas):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = t
        r.font.size = Pt(tam * 0.5); r.font.color.rgb = cor; r.font.name = fonte; r.font.bold = neg
        p.space_after = Pt(tam * 0.22)
    return tb


def topo(sl, s, i, total):
    texto(sl, 120, 60, 1200, 40, [(s.get('k', '').upper(), 22, COBRE, F_MONO, True)])
    texto(sl, 1300, 60, 500, 40, [('MASTERCLASS DE HARDWARE', 22, RGBColor(0x8A, 0x8F, 0x97), F_MONO, True)]).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    texto(sl, 1500, 1010, 300, 40, [(f'{i:02d}/{total:02d}', 22, TINTA, F_MONO, True)]).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    texto(sl, 120, 130, 1600, 120, [(s['t'], 76, TINTA, F_TIT, False)], MSO_ANCHOR.BOTTOM)
    caixa(sl, 120, 262, 1680, 5, TINTA)
    if s.get('rodape'):
        texto(sl, 120, 900, 1600, 90, [(s['rodape'], 27, CINZA, F_TXT, False)], MSO_ANCHOR.BOTTOM)


def gerar_pptx(slides):
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    branco = prs.slide_layouts[6]
    total = len(slides)
    for i, s in enumerate(slides, 1):
        sl = prs.slides.add_slide(branco)
        t = s['tipo']
        if t in ('capa', 'secao', 'fim'):
            fundo(sl, TINTA)
            caixa(sl, 1640, 0, 6, 1080, COBRE)
            if t == 'capa':
                texto(sl, 140, 190, 960, 50, [(s['k'].upper(), 26, COBRE_C, F_MONO, True)])
                texto(sl, 140, 280, 1440, 380, [(s['t'], 150, PAPEL, F_TIT, False)])
                texto(sl, 140, 690, 900, 160, [(s['s'], 36, CINZA_C, F_TXT, False)])
                caixa(sl, 140, 940, 1640, 3, COBRE_C)
                texto(sl, 140, 965, 1000, 40, [('Apostila · HTML interativo · Apresentação', 24, CINZA_C, F_MONO, True)])
                texto(sl, 1300, 965, 480, 40, [(s['r'], 24, CINZA_C, F_MONO, True)]).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
            elif t == 'secao':
                texto(sl, 140, 100, 960, 400, [(s['n'], 300, COBRE, F_TIT, False)])
                texto(sl, 140, 640, 980, 180, [(s['t'], 110, PAPEL, F_TIT, False)], MSO_ANCHOR.BOTTOM)
                texto(sl, 140, 840, 960, 80, [(s['s'], 36, CINZA_C, F_TXT, False)])
                texto(sl, 1500, 1010, 300, 40, [(f'{i:02d}/{total:02d}', 22, PAPEL, F_MONO, True)]).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
            else:
                texto(sl, 140, 300, 980, 260, [(s['t'], 180, PAPEL, F_TIT, False)])
                texto(sl, 140, 600, 900, 80, [(s['s'], 36, CINZA_C, F_TXT, False)])
                texto(sl, 140, 720, 900, 120, [(s['r'], 26, RGBColor(0x9A, 0x9E, 0xA6), F_TXT, False)])
            continue
        fundo(sl, PAPEL)
        topo(sl, s, i, total)
        y0, alt = 310, 560 if s.get('rodape') else 650
        if t == 'cards':
            n = len(s['cards']); cols = 5 if n == 10 else 3
            linhas_n = -(-n // cols); gap = 24
            cw = (1680 - gap * (cols - 1)) / cols; ch = min((alt - gap * (linhas_n - 1)) / linhas_n, 420)
            for j, (a, b, c) in enumerate(s['cards']):
                x = 120 + (j % cols) * (cw + gap); y = y0 + (j // cols) * (ch + gap)
                caixa(sl, x, y, cw, ch, BRANCO, LINHA, True)
                grande = n == 3
                ls = [(a, 44 if n == 10 else 26, COBRE, F_TIT if n == 10 else F_MONO, True), (b, 56 if grande else (34 if n >= 9 else 40), TINTA, F_TIT, False)]
                if c:
                    ls.append((c, 34 if grande else (24 if n >= 9 else 28), CINZA, F_TXT, False))
                texto(sl, x + 26, y + 22, cw - 52, ch - 40, ls)
        elif t == 'duas':
            tem_img = bool(s.get('img'))
            cw = 560 if tem_img else 820
            for j, (lista, escuro) in enumerate(((s['a'], False), (s['b'], True))):
                x = 120 + j * (cw + 40)
                caixa(sl, x, y0, cw, alt, TINTA if escuro else BRANCO, None if escuro else LINHA, True)
                ls = [(lista[0], 56, COBRE_C if escuro else PCB, F_TIT, False), (lista[1], 34, PAPEL if escuro else TINTA, F_TXT, True)]
                ls += [('▪  ' + x_, 31, CINZA_C if escuro else CINZA, F_TXT, False) for x_ in lista[2:]]
                texto(sl, x + 40, y0 + 36, cw - 80, alt - 60, ls)
            if tem_img:
                sl.shapes.add_picture(os.path.join(SAIDA, s['img']), px(1360), px(y0), px(440), px(alt))
        elif t == 'tabela':
            cab, linhas = s['cab'], s['linhas']
            tab = sl.shapes.add_table(len(linhas) + 1, len(cab), px(120), px(y0), px(1680), px(min(alt, 70 * (len(linhas) + 1)))).table
            for c, v in enumerate(cab):
                cel = tab.cell(0, c); cel.text = v or ' '; cel.fill.solid(); cel.fill.fore_color.rgb = TINTA
                r = cel.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = PAPEL; r.font.name = F_TXT
            for l, linha in enumerate(linhas, 1):
                for c, v in enumerate(linha):
                    cel = tab.cell(l, c); cel.text = v; cel.fill.solid()
                    cel.fill.fore_color.rgb = BRANCO if l % 2 else RGBColor(0xFA, 0xF8, 0xF4)
                    r = cel.text_frame.paragraphs[0].runs[0] if cel.text_frame.paragraphs[0].runs else None
                    if r:
                        r.font.size = Pt(15); r.font.color.rgb = TINTA; r.font.name = F_TXT; r.font.bold = c == 0
        elif t == 'numeros':
            nums = s['nums']; tem_img = bool(s.get('img'))
            if tem_img:
                for j, (a, b) in enumerate(nums):
                    x = 120 + (j % 2) * 440; y = y0 + (j // 2) * 190
                    caixa(sl, x, y, 400, 5, COBRE)
                    texto(sl, x, y + 18, 400, 160, [(a, 62, TINTA, F_TIT, False), (b, 26, CINZA, F_TXT, False)])
                sl.shapes.add_picture(os.path.join(SAIDA, s['img']), px(1100), px(y0), px(700), px(430))
                texto(sl, 1100, y0 + 440, 700, 40, [(s.get('legenda', ''), 17, RGBColor(0x8A, 0x8F, 0x97), F_TXT, False)])
            else:
                n = len(nums); gap = 28; cw = (1680 - gap * (n - 1)) / n
                for j, (a, b) in enumerate(nums):
                    x = 120 + j * (cw + gap)
                    caixa(sl, x, y0 + 120, cw, 6, COBRE)
                    texto(sl, x, y0 + 150, cw, 300, [(a, 96 if n <= 3 else (78 if n == 4 else 64), TINTA, F_TIT, False), (b, 28, CINZA, F_TXT, False)])
        elif t == 'linha':
            cw = (1680 - 3 * 32) / 4
            for j, (a, b, c) in enumerate(s['itens']):
                x = 120 + (j % 4) * (cw + 32); y = y0 + (j // 4) * 300
                caixa(sl, x, y, 5, 230, COBRE)
                texto(sl, x + 24, y, cw - 30, 240, [(a, 28, COBRE, F_MONO, True), (b, 42, TINTA, F_TIT, False), (c, 28, CINZA, F_TXT, False)])
        elif t == 'fluxo':
            cw = 560
            for j, (a, b) in enumerate(s['passos']):
                x = 120 + j * cw; escuro = j == 1
                caixa(sl, x, y0 + 40, cw, 380, TINTA if escuro else BRANCO, TINTA)
                texto(sl, x + 44, y0 + 80, cw - 88, 320, [(f'{j + 1:02d}', 28, COBRE, F_MONO, True), (a, 60, PAPEL if escuro else TINTA, F_TIT, False),
                                                          (b, 33, CINZA_C if escuro else CINZA, F_TXT, False)])
    prs.save(os.path.join(SAIDA, 'Masterclass-de-Hardware.pptx'))


def main():
    slides = json.load(io.open(os.path.join(RAIZ, 'conteudo', 'slides.json'), encoding='utf-8'))
    gerar_html(slides)
    gerar_pptx(slides)
    print('ok', len(slides), 'slides')


if __name__ == '__main__':
    main()
