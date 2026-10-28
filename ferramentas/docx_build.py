"""Gera o documento Word (.docx) a partir do HTML dos módulos (saida/modulos_render.json)."""
import io, json, os, re
from html.parser import HTMLParser
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm, Emu

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, 'saida')
TINTA = RGBColor(0x16, 0x18, 0x1C)
COBRE = RGBColor(0xB4, 0x53, 0x1F)
CINZA = RGBColor(0x6E, 0x73, 0x7B)
CORES_CALLOUT = {'exemplo': ('E9F3EE', '0F5A43'), 'dica': ('FAEEE5', 'B4531F'), 'atencao': ('FBEAE7', 'A8261A'),
                 'curiosidade': ('E9F0F9', '1F4E8C'), 'especialista': ('F1EFEA', '16181C')}


# ---------- árvore HTML mínima ----------
class No:
    def __init__(self, tag, attrs=None, pai=None):
        self.tag, self.attrs, self.pai, self.filhos = tag, dict(attrs or []), pai, []

    def classe(self):
        return self.attrs.get('class', '').split()


VAZIOS = {'img', 'br', 'hr', 'meta', 'input'}


class Arvore(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.raiz = No('raiz')
        self.atual = self.raiz

    def handle_starttag(self, tag, attrs):
        no = No(tag, attrs, self.atual)
        self.atual.filhos.append(no)
        if tag not in VAZIOS:
            self.atual = no

    def handle_endtag(self, tag):
        no = self.atual
        while no is not None and no.tag != tag:
            no = no.pai
        if no is not None and no.pai is not None:
            self.atual = no.pai

    def handle_data(self, data):
        self.atual.filhos.append(data)


def arvore(h):
    p = Arvore()
    p.feed(h)
    return p.raiz


def texto(no):
    if isinstance(no, str):
        return no
    return ''.join(texto(f) for f in no.filhos)


# ---------- utilidades docx ----------
def sombrear(celula, cor):
    tcPr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), cor)
    tcPr.append(shd)


def borda_esquerda(celula, cor, tamanho=24):
    tcPr = celula._tc.get_or_add_tcPr()
    bordas = OxmlElement('w:tcBorders')
    for lado in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{lado}')
        if lado == 'left':
            b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), str(tamanho)); b.set(qn('w:color'), cor)
        else:
            b.set(qn('w:val'), 'nil')
        bordas.append(b)
    tcPr.append(bordas)


def cabecalho_repetido(linha):
    trPr = linha._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)


def hyperlink(paragrafo, url, rotulo):
    parte = paragrafo.part
    rid = parte.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '8F3F15'); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    r.append(rPr); t = OxmlElement('w:t'); t.text = rotulo; t.set(qn('xml:space'), 'preserve'); r.append(t)
    h.append(r); paragrafo._p.append(h)


class Escritor:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        estilos = self.doc.styles
        normal = estilos['Normal']
        normal.font.name = 'Calibri'; normal.font.size = Pt(10.5); normal.font.color.rgb = TINTA
        normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
        normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.15
        for nome, tam, cor in (('Title', 34, TINTA), ('Heading 1', 22, TINTA), ('Heading 2', 15, TINTA), ('Heading 3', 12.5, TINTA), ('Heading 4', 11, TINTA)):
            e = estilos[nome]
            e.font.name = 'Arial'; e.font.size = Pt(tam); e.font.bold = True; e.font.color.rgb = cor
            e.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
            e.paragraph_format.space_before = Pt(14 if nome != 'Heading 1' else 0)
            e.paragraph_format.space_after = Pt(6)
            e.paragraph_format.keep_with_next = True
        rodape = sec.footer.paragraphs[0]
        rodape.text = 'Masterclass de Hardware · '
        rodape.runs[0].font.size = Pt(8); rodape.runs[0].font.color.rgb = CINZA
        self._campo_pagina(rodape)

    def _campo_pagina(self, p):
        r = p.add_run(); r.font.size = Pt(8); r.font.color.rgb = COBRE
        for tipo, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
            if tipo:
                f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), tipo); r._r.append(f)
            else:
                it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; r._r.append(it)

    # ---- inline ----
    def inline(self, paragrafo, no, fmt=None):
        fmt = dict(fmt or {})
        if isinstance(no, str):
            t = re.sub(r'\s+', ' ', no)
            if not t:
                return
            r = paragrafo.add_run(t)
            r.bold = fmt.get('b'); r.italic = fmt.get('i')
            if fmt.get('code'):
                r.font.name = 'Consolas'; r.font.size = Pt(9.5)
            if fmt.get('cor'):
                r.font.color.rgb = fmt['cor']
            if fmt.get('tam'):
                r.font.size = fmt['tam']
            return
        tag = no.tag
        if tag in ('strong', 'b'):
            fmt['b'] = True
        elif tag in ('em', 'i'):
            fmt['i'] = True
        elif tag == 'code':
            fmt['code'] = True
        elif tag == 'br':
            paragrafo.add_run().add_break(); return
        elif tag == 'a' and no.attrs.get('href', '').startswith('http'):
            hyperlink(paragrafo, no.attrs['href'], texto(no).strip()); return
        elif tag == 'small':
            fmt['cor'] = CINZA; fmt['tam'] = Pt(8)
        elif tag == 'span' and 'v-m' in no.classe():
            fmt['cor'] = CINZA
            paragrafo.add_run(' — ')
        for f in no.filhos:
            self.inline(paragrafo, f, fmt)

    def paragrafo(self, no, estilo=None, container=None, fmt=None):
        alvo = container or self.doc
        p = alvo.add_paragraph(style=estilo)
        for f in (no.filhos if not isinstance(no, str) else [no]):
            self.inline(p, f, fmt)
        while p.runs and not p.runs[0].text.strip() and not p.runs[0]._r.findall(qn('w:br')):
            p.runs[0]._r.getparent().remove(p.runs[0]._r)
        if p.runs:
            p.runs[0].text = p.runs[0].text.lstrip()
        return p

    # ---- blocos ----
    def blocos(self, no, container=None, nivel_lista=0):
        for f in no.filhos:
            if isinstance(f, str):
                if f.strip():
                    self.paragrafo(f, container=container)
                continue
            self.bloco(f, container, nivel_lista)

    def bloco(self, no, container=None, nivel_lista=0):
        tag, cls = no.tag, no.classe()
        alvo = container or self.doc
        if tag in ('h1', 'h2', 'h3', 'h4'):
            nivel = {'h1': 1, 'h2': 2, 'h3': 3, 'h4': 4}[tag]
            if container is not None:
                p = self.paragrafo(no, container=container, fmt={'b': True})
            else:
                p = self.paragrafo(no, estilo=f'Heading {nivel}')
                if 'divisor' in cls:
                    for r in p.runs:
                        r.font.color.rgb = COBRE; r.font.size = Pt(19)
        elif tag == 'p':
            if 'callout-t' in cls:
                p = alvo.add_paragraph()
                for f in no.filhos:
                    if not isinstance(f, str) and f.tag == 'span':
                        r = p.add_run(texto(f).upper() + '  '); r.bold = True; r.font.size = Pt(8.5)
                        r.font.color.rgb = RGBColor.from_string(self._cor_atual)
                    else:
                        self.inline(p, f, {'b': True})
            else:
                self.paragrafo(no, container=container)
        elif tag in ('ul', 'ol'):
            for i, li in enumerate([x for x in no.filhos if not isinstance(x, str) and x.tag == 'li']):
                if tag == 'ol':
                    # numeração manual: o estilo List Number do Word continua a contagem entre listas
                    p = alvo.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.75 + 0.6 * nivel_lista)
                    p.paragraph_format.first_line_indent = Cm(-0.6)
                    p.add_run(f'{i + 1}.\t').bold = True
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75 + 0.6 * nivel_lista))
                else:
                    p = alvo.add_paragraph(style='List Bullet 2' if nivel_lista else 'List Bullet')
                sub = []
                for f in li.filhos:
                    if not isinstance(f, str) and f.tag in ('ul', 'ol'):
                        sub.append(f)
                    elif not isinstance(f, str) and f.tag == 'p':
                        for g in f.filhos:
                            self.inline(p, g)
                    else:
                        self.inline(p, f)
                if p.runs:
                    p.runs[0].text = p.runs[0].text.lstrip()
                for s in sub:
                    self.bloco(s, container, nivel_lista + 1)
        elif tag == 'div' and 'tabela' in cls:
            for t in no.filhos:
                if not isinstance(t, str) and t.tag == 'table':
                    self.tabela(t, alvo)
        elif tag == 'table':
            self.tabela(no, alvo)
        elif tag == 'div' and 'callout' in cls:
            tipo = next((c[2:] for c in cls if c.startswith('c-')), 'especialista')
            fundo, borda = CORES_CALLOUT.get(tipo, ('F1EFEA', '16181C'))
            self._cor_atual = borda
            t = alvo.add_table(rows=1, cols=1)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            cel = t.cell(0, 0)
            sombrear(cel, fundo); borda_esquerda(cel, borda)
            cel.paragraphs[0]._p.getparent().remove(cel.paragraphs[0]._p)
            self.blocos(no, container=cel)
            if not cel.paragraphs:
                cel.add_paragraph()
            alvo.add_paragraph().paragraph_format.space_after = Pt(2)
        elif tag == 'pre':
            p = alvo.add_paragraph()
            r = p.add_run(texto(no).rstrip('\n')); r.font.name = 'Consolas'; r.font.size = Pt(8)
            p.paragraph_format.line_spacing = 1.0
        elif tag == 'div' and 'galeria' in cls:
            self.galeria(no, alvo)
        elif tag == 'blockquote':
            self.blocos(no, container)
        elif tag in ('section', 'div', 'details'):
            self.blocos(no, container, nivel_lista)
        elif tag == 'hr':
            pass
        else:
            self.blocos(no, container, nivel_lista)

    def tabela(self, no, alvo):
        linhas = []
        for sec in no.filhos:
            if isinstance(sec, str):
                continue
            for tr in ([sec] if sec.tag == 'tr' else sec.filhos):
                if isinstance(tr, str) or tr.tag != 'tr':
                    continue
                celulas = [c for c in tr.filhos if not isinstance(c, str) and c.tag in ('th', 'td')]
                linhas.append((sec.tag == 'thead', celulas))
        if not linhas:
            return
        ncol = max(len(c) for _, c in linhas)
        t = alvo.add_table(rows=len(linhas), cols=ncol)
        t.style = 'Table Grid'
        for i, (cab, celulas) in enumerate(linhas):
            if cab:
                cabecalho_repetido(t.rows[i])
            for j, c in enumerate(celulas):
                cel = t.cell(i, j)
                p = cel.paragraphs[0]
                for f in c.filhos:
                    self.inline(p, f, {'b': True, 'cor': RGBColor(0xF5, 0xF3, 0xEE)} if cab else None)
                for r in p.runs:
                    r.font.size = Pt(8.5)
                if p.runs:
                    p.runs[0].text = p.runs[0].text.lstrip()
                p.paragraph_format.space_after = Pt(0)
                if cab:
                    sombrear(cel, '16181C')
                elif i % 2 == 0:
                    sombrear(cel, 'F6F4EF')
        alvo.add_paragraph().paragraph_format.space_after = Pt(0)

    def galeria(self, no, alvo):
        figuras = [f for f in no.filhos if not isinstance(f, str) and f.tag == 'figure']
        if not figuras:
            return
        cols = 1 if len(figuras) == 1 else 3
        largura = Cm(11) if cols == 1 else Cm(5.2)
        t = alvo.add_table(rows=-(-len(figuras) // cols), cols=cols)
        for k, fig in enumerate(figuras):
            cel = t.cell(k // cols, k % cols)
            img = next(f for f in fig.filhos if not isinstance(f, str) and f.tag == 'img')
            cap = next((f for f in fig.filhos if not isinstance(f, str) and f.tag == 'figcaption'), None)
            p = cel.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(os.path.join(SAIDA, img.attrs['src']), width=largura)
            if cap:
                pc = cel.add_paragraph()
                for f in cap.filhos:
                    if isinstance(f, str):
                        r = pc.add_run(f.strip()); r.bold = True; r.font.size = Pt(8)
                    elif f.tag == 'small':
                        pc.add_run().add_break()
                        r = pc.add_run(texto(f).strip()); r.font.size = Pt(6.5); r.font.color.rgb = CINZA
        alvo.add_paragraph()

    def quebra(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def main():
    modulos = json.load(io.open(os.path.join(SAIDA, 'modulos_render.json'), encoding='utf-8'))
    e = Escritor()
    d = e.doc
    p = d.add_paragraph(); r = p.add_run('MASTERCLASS COMPLETA · TECNOLOGIA DA INFORMAÇÃO'); r.font.size = Pt(9); r.bold = True; r.font.color.rgb = COBRE
    p.paragraph_format.space_before = Pt(120)
    d.add_paragraph('Masterclass de Hardware', style='Title')
    p = d.add_paragraph(); r = p.add_run('Do ábaco à computação quântica: história, fundamentos e cada componente do computador'); r.font.size = Pt(15); r.font.color.rgb = CINZA
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(200)
    r = p.add_run(f'Apostila completa · {len(modulos)} módulos · Outubro de 2026'); r.font.size = Pt(10); r.bold = True
    e.quebra()
    d.add_paragraph('Sumário', style='Heading 1')
    parte = None
    for m in modulos:
        if m['parte'] != parte:
            parte = m['parte']
            p = d.add_paragraph(); r = p.add_run(parte.upper()); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = COBRE
            p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
        p = d.add_paragraph(f'{m["id"]}   {m["titulo"]}'); p.paragraph_format.space_after = Pt(1)
    for m in modulos:
        e.quebra()
        p = d.add_paragraph(); r = p.add_run(f'{m["parte"].upper()} · MÓDULO {m["id"]}'); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = COBRE
        p.paragraph_format.space_after = Pt(2)
        d.add_paragraph(m['titulo'], style='Heading 1')
        e.blocos(arvore(m['figuras']))
        e.blocos(arvore(m['html_p']))
        if m['videos']:
            e.blocos(arvore(m['videos']))
    destino = os.path.join(SAIDA, 'Masterclass-de-Hardware.docx')
    d.save(destino)
    print('ok', destino)


if __name__ == '__main__':
    main()
